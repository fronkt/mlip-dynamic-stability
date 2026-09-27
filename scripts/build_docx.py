"""Build the RSC Advances submission .docx (and optionally .pdf) from the paper markdown.

RSC Advances wants the revised manuscript as a Word file with its figures embedded, a revised
ESI, and figures of at least 600 dpi. The markdown stays the single source of truth: this script
converts paper/manuscript.md and paper/supplementary.md with pandoc and then corrects, with
python-docx, what pandoc's default reference document gets wrong for a submission:

  * pandoc sizes images against its own reference page (at most 14.8 cm wide), so every figure
    is re-placed at the text width, or at its natural size if narrower. Placing a file at or
    below its natural size can only raise its resolution, never lower it.
  * pandoc's styles take their fonts from the document theme (asciiTheme/hAnsiTheme), and a
    theme attribute beats an explicit font name, so the theme attributes are removed rather than
    overridden.
  * captions get Word's built-in Caption style, below the figure and above the table; tables get
    a light grid, a bold repeating header row and 10 pt single-spaced text.
  * A4, 2.5 cm margins, 12 pt Times New Roman, 1.5 line spacing, continuous line numbers and a
    page number in the footer.

The text width is A4 less two 2.5 cm margins, 16.0 cm. RSC's 17.1 cm double-column figure width
would overhang the right margin by a centimetre, so figures are capped at 16.0 cm, and at 22.7 cm
high (the 24.7 cm text height less room for the caption's first lines).

HTML comments never reach the output. A ``<!-- PENDING ... -->`` marker stands for text the paper
still owes, so the build refuses while any remain; ``--allow-pending`` strips them for a draft
build and says so loudly. A '<!--' with no '-->' is refused even then: pandoc would print it.

Run from the repo root:

    python scripts/build_docx.py [SOURCE.md ...] [--out-dir DIR] [--pdf]
                                 [--allow-pending] [--allow-missing-figures] [--check]

With no SOURCE it builds paper/manuscript.md and paper/supplementary.md into paper/, one
<stem>.docx per source. ``--pdf`` also exports <stem>.pdf beside each .docx through Word (COM)
and reports page counts; Word's PDF export re-samples the figures to 200 dpi whatever the file
says, so the PDF is a reading copy and the .docx is what carries the source PNGs byte for byte.
Every build verifies what it wrote (no PENDING text, no '<!--', figure count and caption order,
every 'Table S' and 'Fig.' reference, title equal to the markdown H1, a reference list numbered
1..N, page setup, fonts, figure resolution and size); ``--check`` does it in a temporary
directory and keeps nothing.
``--allow-missing-figures`` exists for comparison builds of old versions such as
paper/submissions/rsc-advances-2026-07/manuscript-as-reviewed.md: a missing figure becomes a
labelled placeholder and figure-quality failures become warnings. It is never right for a file
that is submitted.

Exit codes: 0 ok; 1 verification failed (without --check the files are still written); 2 PENDING
markers or an unclosed comment remain; 3 a figure failed its checks.
"""

from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from PIL import Image, ImageDraw, ImageFont

REPO = Path(__file__).resolve().parents[1]
DEFAULT_SOURCES = [REPO / "paper" / "manuscript.md", REPO / "paper" / "supplementary.md"]
DEFAULT_OUT = REPO / "paper"

PAGE_W_CM, PAGE_H_CM, MARGIN_CM = 21.0, 29.7, 2.5
HEADER_FOOTER_CM = 1.25   # Word's default; pgMar is schema-invalid without header/footer/gutter
TEXT_W_CM = PAGE_W_CM - 2 * MARGIN_CM
# A figure taller than the text block is clipped at the page foot, and keepNext holds the first
# caption lines on its page, so the height is capped below the text height.
MAX_FIG_H_CM = PAGE_H_CM - 2 * MARGIN_CM - 2.0
MIN_DPI = 600
DPI_TOL = 0.5          # PNG stores pixels per metre as an integer: 600 dpi reads back as 599.9988
BODY_FONT, CODE_FONT = "Times New Roman", "Consolas"
BODY_PT, TABLE_PT = 12, 10
LINE_1_5 = "360"       # w:line in 240ths of a line under lineRule="auto"
GRID_COLOR = "BFBFBF"
HEADING_PT = {"Heading1": 16, "Heading2": 14, "Heading3": 12, "Heading4": 12}
TABLE_TEXT = "TableText"
WD_FORMAT_PDF, WD_STATISTIC_PAGES, WD_DO_NOT_SAVE = 17, 2, 0

COMMENT = re.compile(r"<!--(.*?)-->", re.S)
# A comment that fills whole lines; (?!-->) keeps one match from swallowing text between two.
WHOLE_LINE_COMMENT = re.compile(r"^[ \t]*<!--(?:(?!-->).)*-->[ \t]*(?:\r?\n|\Z)", re.S | re.M)
INLINE_COMMENT = re.compile(r"[ \t]*<!--.*?-->", re.S)
# pandoc's smart punctuation turns the "--" of a comment opener that survived into a dash.
COMMENT_LEAK = re.compile(r"<!(?:--|–|—)")
# The alt text may hold bracket pairs ("[0.36, 0.70]", or a link); pandoc needs them balanced too.
IMAGE = re.compile(r"!\[(?P<alt>(?:[^\[\]]|\[[^\[\]]*\])*)\]\((?P<src>[^)\s]+)(?:\s+\"[^\"]*\")?\)"
                   r"(?P<attr>\{[^}]*\})?", re.S)
FIG_LABEL = re.compile(r"^[\s*_]*(Fig\.\s*S?\d+)")
REF_LABEL = re.compile(r"Tables?\s+S\d+|Figs?\.\s*S?\d+")
TABLE_CAPTION = re.compile(r"^\s*Table\s+S?\d+\b")
TABLE_LABEL = re.compile(r"^\s*Table\s+S?\d+\s*$")
H1 = re.compile(r"^#\s+(.+?)\s*#*\s*$", re.M)
LEAKS = {"^": "unconverted superscript caret", "**": "literal bold marker",
         "](": "literal link/image syntax", "$$": "unconverted display math"}
REFS_SECTION = re.compile(r"^#+\s+References\s*\n(.*?)(?=^#|\Z)", re.M | re.S)
REF_ITEM = re.compile(r"^(\d+)\.\s", re.M)
CITATION = re.compile(r"\^(\d[\d,–-]*)\^")
CODE_SPAN = re.compile(r"`[^`]*`")
# pandoc's docx writer drops raw HTML: the text inside survives, the formatting does not.
RAW_HTML = re.compile(r"</?(?:sub|sup|br|span|div|p|b|i|u|em|strong|small|a|img|font)\b[^>]*>",
                      re.I)

M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"
XML_TEXT = re.compile(r"<(?:w|m):t(?:\s[^>]*)?>([^<]*)</(?:w|m):t>")
THEME_ATTRS = [qn(f"w:{a}") for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")]
CODE_STYLE = re.compile(r"^(VerbatimChar|SourceCode|\w+Tok)$")

# Schema order of the children we create; Word rejects a file whose elements are out of order.
RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
             "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
             "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
             "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
             "lang", "eastAsianLayout", "specVanish", "oMath"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
             "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
             "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
STYLE_ORDER = ["name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden",
               "uiPriority", "semiHidden", "unhideWhenUsed", "qFormat", "locked", "personal",
               "personalCompose", "personalReply", "rsid", "pPr", "rPr", "tblPr", "trPr", "tcPr",
               "tblStylePr"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
               "tblPrChange"]
SETTINGS_TAIL = ["doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
                 "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                 "decimalSymbol", "listSeparator"]
SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz",
                "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt",
                "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
                "printerSettings", "sectPrChange"]


@dataclass
class Figure:
    label: str                   # "Fig. 3" from the caption, "" if the caption carries none
    src: str                     # the link as written in the markdown
    path: Path | None            # None when the file is missing and a placeholder stands in
    md5: str = ""
    problems: list[str] = field(default_factory=list)


@dataclass
class Source:
    path: Path
    text: str                    # markdown with every HTML comment removed
    pending: list[tuple[int, str]]
    title: str
    figures: list[Figure] = field(default_factory=list)
    unclosed: list[int] = field(default_factory=list)   # lines of a '<!--' with no '-->'


# ---------------------------------------------------------------- markdown side


def strip_comments(raw: str) -> tuple[str, list[tuple[int, str]]]:
    """Remove every HTML comment and return the text with the (line, body) of each comment.

    A comment that fills its own lines goes with its line break, so a marker between two lines of
    one paragraph (as in the abstract) does not split that paragraph in two.
    """
    found = [(raw.count("\n", 0, m.start()) + 1, " ".join(m.group(1).split()))
             for m in COMMENT.finditer(raw)]
    text = WHOLE_LINE_COMMENT.sub("", raw)
    return INLINE_COMMENT.sub("", text), found


def read_source(path: Path) -> Source:
    """Read and comment-strip one markdown file.

    A '<!--' that is never closed is no comment to pandoc: it would print the marker, PENDING and
    all, into the document as text. Those are recorded in ``unclosed`` for the caller to refuse.
    """
    raw = path.read_text(encoding="utf-8")
    text, comments = strip_comments(raw)
    pending = [(ln, body) for ln, body in comments if body.startswith("PENDING")]
    blanked = COMMENT.sub(lambda m: " " * len(m.group(0)), raw)       # offsets kept
    unclosed = [raw.count("\n", 0, m.start()) + 1 for m in re.finditer("<!--", blanked)]
    m = H1.search(text)
    title = " ".join(m.group(1).replace("*", "").replace("`", "").split()) if m else ""
    return Source(path, text, pending, title, unclosed=unclosed)


def _md5(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()


def check_figures(src: Source, work: Path, lenient: bool) -> list[str]:
    """Check every linked image exists, is RGB and is at least MIN_DPI; return the problems.

    With ``lenient`` a missing file is swapped for a placeholder so a comparison build still has
    one image per caption. Pandoc resolves the remaining links against --resource-path.
    """
    problems: list[str] = []

    def swap(m: re.Match) -> str:
        link = m.group("src")
        lab = FIG_LABEL.match(m.group("alt"))
        fig = Figure(lab.group(1) if lab else "", link, (src.path.parent / link).resolve())
        src.figures.append(fig)
        name = fig.label or link
        if not fig.path.is_file():
            fig.problems.append(f"{src.path.name}: {name}: file not found: {fig.path}")
            fig.path = None
            if not lenient:
                return m.group(0)
            holder = work / f"{src.path.stem}-missing-{len(src.figures)}.png"
            make_placeholder(holder, link)
            return f"![{m.group('alt')}]({holder.name})"
        fig.md5 = _md5(fig.path)
        with Image.open(fig.path) as im:
            mode, dpi = im.mode, im.info.get("dpi")
        if mode != "RGB":
            fig.problems.append(f"{src.path.name}: {name}: {fig.path.name} is {mode}, not RGB")
        if not dpi or min(dpi) < MIN_DPI - DPI_TOL:
            got = f"{min(dpi):.0f} dpi" if dpi else "no dpi metadata"
            fig.problems.append(f"{src.path.name}: {name}: {fig.path.name} is {got}, "
                                f"below {MIN_DPI}")
        return m.group(0)

    src.text = IMAGE.sub(swap, src.text)
    for fig in src.figures:
        problems.extend(fig.problems)
    return problems


def make_placeholder(path: Path, missing: str) -> None:
    w, h = round(TEXT_W_CM / 2.54 * MIN_DPI), round(4.0 / 2.54 * MIN_DPI)
    im = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(im)
    draw.rectangle([6, 6, w - 7, h - 7], outline=(140, 140, 140), width=10)
    try:
        font = ImageFont.load_default(size=90)
    except TypeError:            # Pillow < 10.1 has only the fixed bitmap font
        font = ImageFont.load_default()
    draw.text((w / 2, h / 2), f"Figure file not found: {missing}", fill=(90, 90, 90),
              font=font, anchor="mm")
    im.save(path, dpi=(MIN_DPI, MIN_DPI))


def run_pandoc(src: Source, work: Path, out: Path) -> list[str]:
    """Convert src.text to out; return pandoc's warnings.

    The comment-stripped text goes to a fresh temporary file in ``work``, deleted afterwards, and
    never to <work>/<stem>.md: a caller passing the source's own directory as ``work`` would
    otherwise overwrite the markdown with its PENDING markers stripped.
    """
    fd, name = tempfile.mkstemp(prefix=f"{src.path.stem}-", suffix=".pandoc.md", dir=work)
    md = Path(name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(src.text)
        search = os.pathsep.join([str(src.path.parent), str(work)])
        proc = subprocess.run(
            ["pandoc", str(md), "--from=markdown", "--to=docx", f"--output={out}",
             f"--resource-path={search}"],
            cwd=work, capture_output=True, text=True, encoding="utf-8", errors="replace")
    finally:
        md.unlink(missing_ok=True)
    if proc.returncode:
        raise SystemExit(f"pandoc failed on {src.path.name} ({proc.returncode}):\n{proc.stderr}")
    return [ln for ln in proc.stderr.splitlines() if ln.strip()]


# ---------------------------------------------------------------- docx side


def _child(parent, tag: str, order: list[str]):
    """Return parent's w:<tag> child, creating it at its schema position if absent."""
    el = parent.find(qn(f"w:{tag}"))
    if el is None:
        el = OxmlElement(f"w:{tag}")
        later = {qn(f"w:{t}") for t in order[order.index(tag) + 1:]}
        nxt = next((c for c in parent if c.tag in later), None)
        if nxt is None:
            parent.append(el)
        else:
            nxt.addprevious(el)
    return el


def _set_fonts(rpr, font: str) -> None:
    rfonts = _child(rpr, "rFonts", RPR_ORDER)
    for attr in THEME_ATTRS:
        rfonts.attrib.pop(attr, None)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), font)


def _set_size(rpr, pt: float) -> None:
    for tag in ("sz", "szCs"):
        _child(rpr, tag, RPR_ORDER).set(qn("w:val"), str(round(pt * 2)))


def _drop(parent, *tags: str) -> None:
    for tag in tags:
        for el in parent.findall(qn(f"w:{tag}")):
            parent.remove(el)


def _style(styles, style_id: str):
    return next((s for s in styles.findall(qn("w:style")) if s.get(qn("w:styleId")) == style_id),
                None)


def _style_of(p) -> str:
    ppr = p.find(qn("w:pPr"))
    ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return ps.get(qn("w:val")) if ps is not None else ""


def _text(el, split_scripts: bool = False) -> str:
    """Text of el. ``split_scripts`` puts a space before super/subscript runs, so 'Fig. 1' with a
    citation after it does not read as 'Fig. 14' when references are counted."""
    parts = []
    for t in el.iter(qn("w:t"), M_T):
        if split_scripts and t.getparent().find(f"{qn('w:rPr')}/{qn('w:vertAlign')}") is not None:
            parts.append(" ")
        parts.append(t.text or "")
    return "".join(parts)


def set_styles(doc) -> None:
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    rpr = _child(_child(defaults, "rPrDefault", ["rPrDefault", "pPrDefault"]), "rPr", ["rPr"])
    _set_fonts(rpr, BODY_FONT)
    _set_size(rpr, BODY_PT)
    ppr = _child(_child(defaults, "pPrDefault", ["rPrDefault", "pPrDefault"]), "pPr", ["pPr"])
    spacing = _child(ppr, "spacing", PPR_ORDER)
    spacing.set(qn("w:line"), LINE_1_5)
    spacing.set(qn("w:lineRule"), "auto")

    for style in styles.findall(qn("w:style")):
        srpr = style.find(qn("w:rPr"))
        if srpr is not None and srpr.find(qn("w:rFonts")) is not None:
            code = CODE_STYLE.match(style.get(qn("w:styleId")) or "")
            _set_fonts(srpr, CODE_FONT if code else BODY_FONT)

    for style_id, pt in HEADING_PT.items():
        style = _style(styles, style_id)
        if style is None:
            continue
        hrpr = _child(style, "rPr", STYLE_ORDER)
        _set_fonts(hrpr, BODY_FONT)
        _drop(hrpr, "color", "i", "iCs")
        _child(hrpr, "b", RPR_ORDER)
        _child(hrpr, "bCs", RPR_ORDER)
        if style_id == "Heading4":
            _child(hrpr, "i", RPR_ORDER)
            _child(hrpr, "iCs", RPR_ORDER)
        _set_size(hrpr, pt)

    caption = _style(styles, "Caption")
    if caption is not None and caption.find(qn("w:rPr")) is not None:
        _drop(caption.find(qn("w:rPr")), "i", "iCs")

    table_text = doc.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text.element.set(qn("w:styleId"), TABLE_TEXT)
    table_text.base_style = doc.styles["Normal"]
    table_text.font.size = Pt(TABLE_PT)
    fmt = table_text.paragraph_format
    fmt.line_spacing, fmt.space_before, fmt.space_after = 1.0, Pt(0), Pt(0)


def set_page(doc) -> None:
    for section in doc.sections:
        section.page_width, section.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
        section.left_margin = section.right_margin = Cm(MARGIN_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_CM)
        section.header_distance = section.footer_distance = Cm(HEADER_FOOTER_CM)
        section.gutter = Cm(0)
        lnnum = _child(section._sectPr, "lnNumType", SECTPR_ORDER)
        lnnum.set(qn("w:countBy"), "1")
        lnnum.set(qn("w:restart"), "continuous")

        footer = section.footer
        footer.is_linked_to_previous = False
        para = footer.paragraphs[0]
        for run in list(para._p.findall(qn("w:r"))):
            para._p.remove(run)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, value in (("fldChar", "begin"), ("instrText", " PAGE "),
                            ("fldChar", "separate"), ("t", "1"), ("fldChar", "end")):
            run, el = OxmlElement("w:r"), OxmlElement(f"w:{kind}")
            if kind == "fldChar":
                el.set(qn("w:fldCharType"), value)
            else:
                el.set(qn("xml:space"), "preserve")
                el.text = value
            run.append(el)
            para._p.append(run)


def place_figures(doc) -> None:
    """Size each figure to its natural width within the text width and MAX_FIG_H_CM height,
    centre it, caption it below."""
    limit, max_h = Cm(TEXT_W_CM), Cm(MAX_FIG_H_CM)
    for shape in doc.inline_shapes:
        blob = doc.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
        with Image.open(io.BytesIO(blob.blob)) as im:
            (px_w, px_h), dpi = im.size, (im.info.get("dpi") or (96, 96))[0]
        width = min(Emu(round(px_w / dpi * 914400)), limit, Emu(round(max_h * px_w / px_h)))
        shape.width, shape.height = Emu(width), Emu(round(width * px_h / px_w))

    body = doc.element.body
    for p in body.findall(qn("w:p")):
        if p.find(".//" + qn("w:drawing")) is None:
            continue
        ppr = p.get_or_add_pPr()
        ppr.keepNext_val = True
        ppr.jc_val = WD_ALIGN_PARAGRAPH.CENTER
        cap = p.getnext()
        if cap is not None and cap.tag == qn("w:p") and _style_of(cap) in ("ImageCaption",
                                                                          "Caption"):
            cap.get_or_add_pPr().style = "Caption"


def _is_table_caption(p) -> bool:
    """A paragraph that opens with a bold 'Table N' / 'Table SN', the papers' caption convention.

    Adjacency to the table is not required: Table S11's caption is followed by two paragraphs of
    notes before its table.
    """
    label = ""
    for r in p.findall(qn("w:r")):
        if r.find(f"{qn('w:rPr')}/{qn('w:b')}") is None:
            break
        label += _text(r)
    return bool(TABLE_LABEL.match(label))


def _table_caption(tbl) -> tuple[str | None, bool]:
    """(text of the 'Table N' caption that introduces tbl or None, whether tbl is a later part).

    Scanning back through paragraphs and reaching another table first makes tbl a later part of
    that table, inheriting its caption: Tables S16 and S17 come in parts under one caption.
    """
    el = tbl.getprevious()
    while (el is not None and el.tag in (qn("w:p"), qn("w:tbl"))
           and not _style_of(el).startswith("Heading")):
        if el.tag == qn("w:tbl"):
            return _table_caption(el)[0], True
        if _style_of(el) == "Caption" and TABLE_CAPTION.match(_text(el)):
            return _text(el), False
        el = el.getprevious()
    return None, False


def format_tables(doc) -> None:
    """Light grid, full text width with Word's autofit columns, bold header, 10 pt text."""
    body = doc.element.body
    for p in body.findall(qn("w:p")):
        if _is_table_caption(p):
            ppr = p.get_or_add_pPr()
            ppr.style = "Caption"
            ppr.keepNext_val = True

    for tbl in body.iter(qn("w:tbl")):
        tblpr = tbl.find(qn("w:tblPr"))
        width = _child(tblpr, "tblW", TBLPR_ORDER)
        width.set(qn("w:type"), "pct")
        width.set(qn("w:w"), "5000")
        _drop(tblpr, "tblLayout")
        borders = _child(tblpr, "tblBorders", TBLPR_ORDER)
        for edge in list(borders):
            borders.remove(edge)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            for key, value in (("val", "single"), ("sz", "4"), ("space", "0"),
                               ("color", GRID_COLOR)):
                el.set(qn(f"w:{key}"), value)
            borders.append(el)

        for tr in tbl.findall(qn("w:tr")):
            header = tr.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None
            for p in tr.iter(qn("w:p")):
                p.get_or_add_pPr().style = TABLE_TEXT
                for r in p.iter(qn("w:r")):
                    rpr = r.get_or_add_rPr()
                    if header:
                        _child(rpr, "b", RPR_ORDER)
                        _child(rpr, "bCs", RPR_ORDER)
                    rstyle = rpr.find(qn("w:rStyle"))
                    if rstyle is not None and rstyle.get(qn("w:val")) == "VerbatimChar":
                        _set_size(rpr, TABLE_PT - 1)   # the code style's 11 pt dwarfs 10 pt text


def keep_full_resolution(doc) -> None:
    """Set Word's 'Do not compress images in file'.

    A Word whose default image resolution is not 'High fidelity' re-samples pictures when it saves
    a file, which would take the 600 dpi figures below RSC's floor the first time anyone re-saves
    the manuscript. (This machine's Word keeps them either way; the editorial office's may not.)
    It does not govern Word's PDF export, which writes the figures at 200 dpi regardless.
    """
    _child(doc.settings.element, "doNotAutoCompressPictures", SETTINGS_TAIL)


def format_docx(path: Path, title: str) -> None:
    doc = Document(str(path))
    set_styles(doc)
    set_page(doc)
    place_figures(doc)
    format_tables(doc)
    keep_full_resolution(doc)
    doc.core_properties.title = title
    doc.save(str(path))


# ---------------------------------------------------------------- inspection and checks


def inspect(path: Path) -> dict:
    """Read a built .docx back and measure what the report and --check need."""
    doc = Document(str(path))
    body = doc.element.body
    images = []
    for shape in doc.inline_shapes:
        blob = doc.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
        with Image.open(io.BytesIO(blob.blob)) as im:
            size, mode, dpi = im.size, im.mode, im.info.get("dpi")
        images.append(dict(px=size, mode=mode, file_dpi=dpi and round(min(dpi)),
                           cm=(shape.width / 360000, shape.height / 360000),
                           placed_dpi=size[0] / (shape.width / 914400),
                           md5=hashlib.md5(blob.blob).hexdigest()))

    paras = body.findall(qn("w:p"))
    captions = [_text(p) for p in paras if _style_of(p) == "Caption"]
    tables = body.findall(qn("w:tbl"))
    table_captions = [_table_caption(t) for t in tables]
    uncaptioned = [_text(t.find(f".//{qn('w:tr')}"))[:60]
                   for t, (cap, _) in zip(tables, table_captions) if cap is None]
    table_parts = sum(1 for cap, part in table_captions if cap is not None and part)
    heading1 = next((_text(p) for p in paras if _style_of(p) == "Heading1"), "")

    section = doc.sections[0]
    lnnum = section._sectPr.find(qn("w:lnNumType"))
    footer_xml = section.footer._element.xml if not section.footer.is_linked_to_previous else ""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rfonts = defaults.find(f".//{qn('w:rFonts')}")
    size = defaults.find(f".//{qn('w:sz')}")
    spacing = defaults.find(f".//{qn('w:spacing')}")

    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n).decode("utf-8", "replace") for n in z.namelist()
                 if n.endswith(".xml")}
    part_text = "\n".join(html.unescape(t) for n, x in parts.items() if n.startswith("word/")
                          for t in XML_TEXT.findall(x))
    theme_left = sum(len(re.findall(r'w:(?:asciiTheme|hAnsiTheme|eastAsiaTheme|cstheme)=', x))
                     for n, x in parts.items() if n in ("word/styles.xml", "word/document.xml"))

    return dict(
        paragraphs=len(doc.paragraphs),
        paragraphs_all=len(body.findall(".//" + qn("w:p"))),
        images=images, captions=captions, tables=len(tables), uncaptioned=uncaptioned,
        table_parts=table_parts,
        title=heading1,
        text="\n".join(_text(p, split_scripts=True) for p in body.iter(qn("w:p"))),
        all_text=part_text,
        superscripts=len(body.findall(".//" + qn("w:vertAlign") + "[@" + qn("w:val")
                                      + "='superscript']")),
        page_cm=(section.page_width.cm, section.page_height.cm),
        margins_cm=tuple(round(m.cm, 2) for m in (section.left_margin, section.right_margin,
                                                  section.top_margin, section.bottom_margin)),
        lnnum=dict((k.split("}")[1], v) for k, v in lnnum.attrib.items()) if lnnum is not None
        else None,
        page_field="PAGE" in footer_xml and "fldChar" in footer_xml,
        fonts=dict((k.split("}")[1], v) for k, v in rfonts.attrib.items()) if rfonts is not None
        else {},
        body_pt=int(size.get(qn("w:val"))) / 2 if size is not None else None,
        line=spacing.get(qn("w:line")) if spacing is not None else None,
        theme_attrs_left=theme_left,
    )


def _norm(s: str) -> str:
    s = s.translate(str.maketrans("“”‘’", "\"\"''"))
    return re.sub(r"\s+", " ", s).strip()


def verify(src: Source, st: dict) -> tuple[list[str], list[str]]:
    """Return (failures, warnings) for one built file against its markdown."""
    fails, warns = [], []
    if "PENDING" in st["all_text"]:
        fails.append("'PENDING' text present in the .docx")
    if COMMENT_LEAK.search(st["all_text"]):
        fails.append("'<!--' (or its smart-dash form '<!–') present in the .docx")

    n_md = len(src.figures)
    if len(st["images"]) != n_md:
        fails.append(f"{len(st['images'])} images in the .docx, {n_md} in the markdown")
    fig_caps = [c for c in st["captions"] if c.lstrip().startswith("Fig")]
    if len(fig_caps) != n_md:
        fails.append(f"{len(fig_caps)} figure captions in the .docx, {n_md} in the markdown")
    numbers = [re.match(r"\s*Fig\.\s*S?(\d+)", c) for c in fig_caps]
    order = [int(m.group(1)) for m in numbers if m]
    if order != list(range(1, len(order) + 1)):
        fails.append(f"figure captions out of order: {order}")
    for i, img in enumerate(st["images"], 1):
        if img["mode"] != "RGB" or img["placed_dpi"] < MIN_DPI - DPI_TOL:
            fails.append(f"image {i}: {img['mode']}, {img['placed_dpi']:.0f} dpi as placed")
        if img["cm"][0] > TEXT_W_CM + 0.01 or img["cm"][1] > MAX_FIG_H_CM + 0.01:
            fails.append(f"image {i}: {img['cm'][0]:.2f} x {img['cm'][1]:.2f} cm, over "
                         f"{TEXT_W_CM} x {MAX_FIG_H_CM} cm")

    want = collections.Counter(_norm(m) for m in REF_LABEL.findall(_norm(src.text)))
    got = collections.Counter(_norm(m) for m in REF_LABEL.findall(_norm(st["text"])))
    lost = want - got
    if lost:
        fails.append("references lost: " + ", ".join(f"{k} x{v}" for k, v in sorted(lost.items())))
    if _norm(st["title"]) != _norm(src.title):
        fails.append(f"title {st['title']!r} != markdown H1 {src.title!r}")

    if tuple(round(x, 2) for x in st["page_cm"]) != (PAGE_W_CM, PAGE_H_CM):
        fails.append(f"page {st['page_cm']} cm, not A4")
    if any(abs(m - MARGIN_CM) > 0.01 for m in st["margins_cm"]):
        fails.append(f"margins {st['margins_cm']} cm")
    if not st["lnnum"] or st["lnnum"].get("restart") != "continuous":
        fails.append(f"line numbering {st['lnnum']}")
    if not st["page_field"]:
        fails.append("no PAGE field in the footer")
    if set(st["fonts"].values()) != {BODY_FONT} or st["theme_attrs_left"]:
        fails.append(f"fonts {st['fonts']}, theme-font attributes left {st['theme_attrs_left']}")

    ref_fails, ref_warns = check_references(src.text)
    fails += ref_fails
    warns += ref_warns
    tags = sorted({m.group(0) for m in RAW_HTML.finditer(CODE_SPAN.sub("", src.text))})
    if tags:
        warns.append(f"raw HTML {', '.join(tags[:6])} in the markdown: pandoc's docx writer drops "
                     "it, so the text inside loses that formatting")
    for token, what in LEAKS.items():
        hits = [ln for ln in st["text"].splitlines() if token in ln]
        if hits:
            warns.append(f"{what} ({token!r}) in {len(hits)} paragraph(s), first: "
                         f"{_snippet(hits[0], token)}")
    for header in st["uncaptioned"]:
        warns.append(f"table without a 'Table N' caption, header row: {header!r}")
    return fails, warns


def _cited(text: str) -> set[int]:
    """Reference numbers cited as ^n^, ^n,m^ or ^n–m^ (en dash or hyphen)."""
    nums: set[int] = set()
    for m in CITATION.finditer(text):
        for part in m.group(1).replace("–", "-").split(","):
            lo, _, hi = part.partition("-")
            if lo.isdigit():
                nums.update(range(int(lo), int(hi) + 1) if hi.isdigit() else {int(lo)})
    return nums


def check_references(text: str) -> tuple[list[str], list[str]]:
    """Return (failures, warnings) for the reference list against the citations.

    pandoc numbers an ordered list from its first item and ignores the rest, so a list written
    1, 2, 4 prints as 1, 2, 3 and every later entry drifts from its citations without a word.
    A superscript that is an exponent rather than a citation cannot be told apart, so citations
    missing from the list only warn. A file with citations and no list (an ESI citing a paper
    the main text does not) is flagged for the same reason.
    """
    fails, warns = [], []
    body = CODE_SPAN.sub("", text)
    section = REFS_SECTION.search(body)
    listed = [int(n) for n in REF_ITEM.findall(section.group(1))] if section else []
    cited = _cited(body[:section.start()] + body[section.end():] if section else body)
    if len(set(listed)) > 1 and listed != list(range(1, len(listed) + 1)):
        wrong = [f"{n} prints as {i}" for i, n in enumerate(listed, 1) if n != i]
        fails.append(f"reference list not numbered 1..{len(listed)}: {', '.join(wrong[:4])}"
                     + (" ..." if len(wrong) > 4 else ""))
    missing = sorted(cited - set(range(1, len(listed) + 1)))
    if missing and section:
        warns.append(f"cited but not in the {len(listed)}-entry reference list: "
                     + ", ".join(map(str, missing[:12])))
    elif missing:
        warns.append("cites reference number(s) " + ", ".join(map(str, missing[:12]))
                     + " but has no References list; check each resolves in the main text's list")
    return fails, warns


def _snippet(line: str, token: str, width: int = 50) -> str:
    i = line.index(token)
    return "..." + line[max(0, i - width):i + width].replace("\n", " ") + "..."


# ---------------------------------------------------------------- Word (PDF export)


def _alive(pid: int) -> bool:
    """True while pid runs. tasklist keeps listing a Word that has exited for as long as anyone
    holds its process handle, so exit is judged by the exit code instead."""
    import ctypes

    k32 = ctypes.windll.kernel32
    handle = k32.OpenProcess(0x1000, False, pid)        # PROCESS_QUERY_LIMITED_INFORMATION
    if not handle:
        return False
    try:
        code = ctypes.c_ulong()
        return bool(k32.GetExitCodeProcess(handle, ctypes.byref(code))) and code.value == 259
    finally:
        k32.CloseHandle(handle)


def _winword_pids() -> set[int]:
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq WINWORD.EXE", "/FO", "CSV", "/NH"],
                         capture_output=True, text=True).stdout
    return {int(row[1]) for row in csv.reader(out.splitlines())
            if len(row) > 1 and row[0].upper() == "WINWORD.EXE" and _alive(int(row[1]))}


def _kill(pids: set[int]) -> None:
    for pid in pids:
        subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)


def _pdf_pages(pdf: Path) -> int | None:
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    return len(PdfReader(str(pdf)).pages)


def export_pdfs(docx_paths: list[Path], timeout: float = 600.0, quit_wait: float = 90.0) -> dict:
    """Export each .docx to <stem>.pdf beside it with Word; return pages and process bookkeeping.

    Word runs as a fresh invisible instance (DispatchEx, never one the user has open), opens each
    file read-only and is always told to Quit. Our process is identified exactly, from the window
    handle of the first document, because other sessions may be driving Word at the same time;
    until that document is open, a single new WINWORD.EXE stands in for it. Word can take half a
    minute to exit after Quit returns, so it gets ``quit_wait`` seconds before it is killed; the
    watchdog kills it if the export itself hangs past ``timeout``.
    """
    import pythoncom
    import win32com.client
    import win32process

    before = _winword_pids()
    ours: set[int] = set()
    new: set[int] = set()
    pages: dict[Path, tuple[int, int | None]] = {}
    expired = threading.Event()

    def expire() -> None:
        expired.set()
        _kill(ours)

    watchdog = threading.Timer(timeout, expire)
    watchdog.daemon = True
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        new = _winword_pids() - before
        ours = new if len(new) == 1 else set()
        watchdog.start()
        word.Visible = False
        word.DisplayAlerts = 0
        for path in docx_paths:
            pdf = path.with_suffix(".pdf")
            doc = word.Documents.Open(str(path), False, True, False)  # confirm, read-only, recent
            try:
                if path == docx_paths[0]:
                    ours = {win32process.GetWindowThreadProcessId(doc.ActiveWindow.Hwnd)[1]}
                doc.Repaginate()
                n_word = doc.ComputeStatistics(WD_STATISTIC_PAGES)
                doc.SaveAs2(str(pdf), WD_FORMAT_PDF)
            finally:
                try:
                    doc.Close(WD_DO_NOT_SAVE)
                except Exception:      # Word died under us; the error that killed it is the one
                    pass               # to report, not Close's AttributeError on a dead server
                doc = None
            pages[path] = (n_word, _pdf_pages(pdf))
    finally:
        dispatched = word is not None
        if dispatched:
            try:
                word.Quit(WD_DO_NOT_SAVE)
            except Exception:          # already dead (watchdog) or disconnected
                pass
            word = None
        watchdog.cancel()
        pythoncom.CoUninitialize()
        start = time.monotonic()
        while any(_alive(p) for p in ours) and time.monotonic() - start < quit_wait:
            time.sleep(0.5)
        exit_s = time.monotonic() - start
        killed = {p for p in ours if _alive(p)}
        _kill(killed)
        if dispatched and not ours:
            print(f"WARNING: could not identify our WINWORD.EXE ({len(new)} appeared); nothing "
                  "was force-killed, check Task Manager", file=sys.stderr)
        if expired.is_set():
            print(f"ERROR: the Word export ran past {timeout:g} s and Word was killed; no PDF "
                  "after that point is complete", file=sys.stderr)
    return dict(pages=pages, before=before, ours=ours, exit_s=exit_s, killed=killed,
                after=_winword_pids())


# ---------------------------------------------------------------- driver


def _banner(lines: list[str]) -> str:
    bar = "!" * 78
    return "\n".join([bar, *("!! " + ln for ln in lines), bar])


def report(src: Source, out: Path, st: dict) -> None:
    print(f"{src.path.name} -> {out}")
    print(f"  paragraphs : {st['paragraphs']} body-level ({st['paragraphs_all']} with table "
          "cells)")
    n_parts, n_bare = st["table_parts"], len(st["uncaptioned"])
    print(f"  tables     : {st['tables']} ({st['tables'] - n_parts - n_bare} under their own "
          f"'Table N' caption, {n_parts} later parts of a captioned table, {n_bare} uncaptioned)")
    print(f"  figures    : {len(st['images'])} embedded, {len(src.figures)} in the markdown")
    for fig, img in zip(src.figures, st["images"]):
        same = "source bytes" if fig.md5 and fig.md5 == img["md5"] else "PLACEHOLDER/CHANGED"
        print(f"    {fig.label or '?':8} {Path(fig.src).name:28} {img['px'][0]}x{img['px'][1]} px "
              f"{img['mode']}, file {img['file_dpi']} dpi, placed {img['cm'][0]:.2f} x "
              f"{img['cm'][1]:.2f} cm = {img['placed_dpi']:.0f} dpi ({same})")
    print(f"  captions   : {len(st['captions'])} Word 'Caption' paragraphs")
    print(f"  superscript: {st['superscripts']} runs")
    print(f"  page       : {st['page_cm'][0]:.1f} x {st['page_cm'][1]:.1f} cm, margins "
          f"{st['margins_cm']} cm, line numbers {st['lnnum']}, footer PAGE field "
          f"{'yes' if st['page_field'] else 'NO'}")
    print(f"  fonts      : defaults {st['fonts']}, {st['body_pt']} pt, line {st['line']}/240; "
          f"theme-font attributes left {st['theme_attrs_left']}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("sources", nargs="*", type=Path, help="markdown files (default: manuscript "
                    "and ESI)")
    ap.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--pdf", action="store_true", help="also export PDF through Word")
    ap.add_argument("--allow-pending", action="store_true",
                    help="strip PENDING markers instead of refusing (draft builds only)")
    ap.add_argument("--allow-missing-figures", action="store_true",
                    help="placeholder for missing figures, figure checks as warnings "
                         "(comparison builds only)")
    ap.add_argument("--check", action="store_true",
                    help="build into a temporary directory, verify, keep nothing")
    args = ap.parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(errors="replace")

    sources = [read_source(p.resolve()) for p in (args.sources or DEFAULT_SOURCES)]
    unclosed = [f"{s.path.name}:{ln}" for s in sources for ln in s.unclosed]
    if unclosed:
        print(f"REFUSED: '<!--' never closed by '-->' at {', '.join(unclosed)}; pandoc would print "
              "the marker into the document as text (--allow-pending cannot strip it)",
              file=sys.stderr)
        return 2
    pending = [(s.path.name, ln, body) for s in sources for ln, body in s.pending]
    listing = [f"{name}:{ln}  {body[:110]}" for name, ln, body in pending]
    if pending and not args.allow_pending:
        print(f"REFUSED: {len(pending)} PENDING marker(s) remain; the text they stand for is not "
              "written yet (--allow-pending strips them for a draft build):", file=sys.stderr)
        print("\n".join("  " + ln for ln in listing), file=sys.stderr)
        return 2
    if pending:
        print(_banner([f"--allow-pending: {len(pending)} PENDING marker(s) STRIPPED. The text "
                       "they stand for is NOT", "in these files. This is a DRAFT build: do NOT "
                       "submit it.", *listing]), file=sys.stderr)

    with tempfile.TemporaryDirectory(prefix="build_docx_") as tmp:
        work = Path(tmp)
        problems = [p for s in sources for p in check_figures(s, work, args.allow_missing_figures)]
        if problems and not args.allow_missing_figures:
            print("REFUSED: figure check failed:", file=sys.stderr)
            print("\n".join("  " + p for p in problems), file=sys.stderr)
            return 3
        for p in problems:
            print(f"WARNING (--allow-missing-figures): {p}", file=sys.stderr)

        out_dir = work / "out" if args.check else args.out_dir.resolve()
        out_dir.mkdir(parents=True, exist_ok=True)
        built, failed = [], False
        for src in sources:
            staged = work / f"{src.path.stem}.docx"
            for warning in run_pandoc(src, work, staged):
                print(f"pandoc: {warning}", file=sys.stderr)
            format_docx(staged, src.title)
            out = out_dir / staged.name
            shutil.copyfile(staged, out)
            st = inspect(out)
            report(src, out, st)
            fails, warns = verify(src, st)
            for f in fails:
                print(f"  FAIL  {f}")
            for w in warns:
                print(f"  WARN  {w}")
            print(f"  check      : {'FAILED' if fails else 'passed'}"
                  + (" (the file is written, but do not submit it)"
                     if fails and not args.check else ""))
            failed = failed or bool(fails)
            built.append(out)

        if args.pdf:
            word = export_pdfs(built)
            for path, (n_word, n_pdf) in word["pages"].items():
                print(f"{path.with_suffix('.pdf')}: {n_word} pages (Word), {n_pdf} pages (PDF)")
            fate = (f"force-killed {sorted(word['killed'])}" if word["killed"]
                    else f"exited {word['exit_s']:.1f} s after Quit")
            print(f"WINWORD.EXE running: before {sorted(word['before'])}, ours "
                  f"{sorted(word['ours'])} ({fate}), after {sorted(word['after'])}")
        if pending:
            print(_banner([f"DRAFT: {len(pending)} PENDING marker(s) were stripped from these "
                           "files."]), file=sys.stderr)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
